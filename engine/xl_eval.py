"""
Small Excel formula evaluator for workbooks written by build_model.py.

Used by verify_model.py when the `formulas` package is not installed. It covers
only the operators and functions the engine writes; anything else raises
NotImplementedError naming the function, so a gap shows up as an error rather
than a wrong number.

It reads the formulas the workbook actually contains (not the Python that wrote
them), so a cell-reference typo in the built file shows up in the results.
"""

from __future__ import annotations

import math
import re
import sys
import threading
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter

EXCEL_EPOCH = date(1899, 12, 30)
CELL_RE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)$")
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


class XLError:
    def __init__(self, code):
        self.code = code

    def __repr__(self):
        return self.code


# ---------------------------------------------------------------------------
# Parser: openpyxl tokens -> nested tuples
# ---------------------------------------------------------------------------
INFIX_PREC = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1,
              "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}
PREFIX_PREC = 6


def parse(formula: str):
    toks = [t for t in Tokenizer(formula).items if t.type != Token.WSPACE]
    pos = 0

    def peek():
        return toks[pos] if pos < len(toks) else None

    def take():
        nonlocal pos
        pos += 1
        return toks[pos - 1]

    def expr(min_prec):
        left = primary()
        while True:
            t = peek()
            if t is None or t.type != Token.OP_IN:
                return left
            prec = INFIX_PREC[t.value]
            if prec < min_prec:
                return left
            take()
            left = ("bin", t.value, left, expr(prec + 1))

    def primary():
        t = take()
        if t.type == Token.OP_PRE:
            node = ("neg", expr(PREFIX_PREC)) if t.value == "-" else expr(PREFIX_PREC)
        elif t.type == Token.OPERAND:
            if t.subtype == Token.NUMBER:
                node = ("num", float(t.value))
            elif t.subtype == Token.TEXT:
                node = ("str", t.value[1:-1].replace('""', '"'))
            elif t.subtype == Token.LOGICAL:
                node = ("num", 1.0 if t.value.upper() == "TRUE" else 0.0)
            elif t.subtype == Token.ERROR:
                node = ("err", t.value)
            else:
                node = ("ref", t.value)
        elif t.type == Token.FUNC and t.subtype == Token.OPEN:
            name, args = t.value[:-1].upper(), []
            if not (peek().type == Token.FUNC and peek().subtype == Token.CLOSE):
                while True:
                    args.append(expr(0))
                    s = take()
                    if s.type == Token.FUNC and s.subtype == Token.CLOSE:
                        break
            else:
                take()
            node = ("func", name, args)
        elif t.type == Token.PAREN and t.subtype == Token.OPEN:
            node = expr(0)
            take()  # closing paren
        else:
            raise ValueError(f"unexpected token {t.value!r} in {formula!r}")
        while peek() is not None and peek().type == Token.OP_POST:
            take()
            node = ("pct", node)
        return node

    tree = expr(0)
    if pos != len(toks):
        raise ValueError(f"trailing tokens in {formula!r}")
    return tree


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _num(v):
    if isinstance(v, XLError):
        raise _Propagate(v)
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        raise _Propagate(XLError("#VALUE!"))


class _Propagate(Exception):
    def __init__(self, err):
        self.err = err


def _flatten(vals):
    for v in vals:
        if isinstance(v, list):
            yield from _flatten(v)
        else:
            yield v


def to_date(serial):
    return EXCEL_EPOCH + timedelta(days=int(serial))


def to_serial(d):
    if isinstance(d, datetime):
        d = d.date()
    return float((d - EXCEL_EPOCH).days)


def _edate(serial, months):
    d = to_date(serial)
    m = d.month - 1 + int(months)
    y, m = d.year + m // 12, m % 12 + 1
    last = [31, 29 if (y % 4 == 0 and (y % 100 != 0 or y % 400 == 0)) else 28,
            31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return to_serial(date(y, m, min(d.day, last)))


def _text(v, fmt):
    f = fmt.lower()
    if f == "mmm yyyy":
        d = to_date(_num(v))
        return f"{MONTHS[d.month - 1]} {d.year}"
    if f == "mmm yy":
        d = to_date(_num(v))
        return f"{MONTHS[d.month - 1]} {d.year % 100:02d}"
    if f == "yyyy":
        return str(to_date(_num(v)).year)
    raise NotImplementedError(f"TEXT format {fmt!r}")


def _roundup(x, digits):
    f = 10 ** int(digits)
    return math.copysign(math.ceil(abs(x) * f - 1e-12) / f, x)


# ---------------------------------------------------------------------------
# Evaluator
# ---------------------------------------------------------------------------
class Evaluator:
    def __init__(self, path):
        self.wb = load_workbook(path)
        self.names = {}
        for name, dn in self.wb.defined_names.items():
            sheet, coord = dn.attr_text.split("!")
            self.names[name] = (sheet.strip("'"), coord.replace("$", ""))
        self.cache = {}
        self.trees = {}

    # public -----------------------------------------------------------------
    def values(self, cells):
        """Evaluate a list of (sheet, coord). Runs in a thread with a deep
        stack, because period-to-period links make long dependency chains."""
        out = {}
        err = []

        def run():
            try:
                for sheet, coord in cells:
                    out[(sheet, coord)] = self.cell(sheet, coord)
            except BaseException as e:  # surface in caller
                err.append(e)

        sys.setrecursionlimit(1_000_000)
        threading.stack_size(512 * 1024 * 1024)
        t = threading.Thread(target=run)
        t.start()
        t.join()
        if err:
            raise err[0]
        return out

    # internals --------------------------------------------------------------
    def cell(self, sheet, coord):
        key = (sheet, coord)
        if key in self.cache:
            return self.cache[key]
        raw = self.wb[sheet][coord].value
        if isinstance(raw, str) and raw.startswith("="):
            self.cache[key] = XLError("#CIRC!")   # seen again before finishing = cycle
            if key not in self.trees:
                self.trees[key] = parse(raw)
            try:
                v = self.ev(self.trees[key], sheet)
            except _Propagate as p:
                v = p.err
            if isinstance(v, list):
                v = v[0]
        elif isinstance(raw, (date, datetime)):
            v = to_serial(raw)
        else:
            v = raw
        self.cache[key] = v
        return v

    def ref(self, text, sheet):
        if "!" in text:
            sheet, text = text.rsplit("!", 1)
            sheet = sheet.strip("'")
        text = text.replace("$", "")
        if ":" in text:
            a, b = text.split(":")
            ca, ra = CELL_RE.match(a).groups()
            cb, rb = CELL_RE.match(b).groups()
            c0, c1 = column_index_from_string(ca), column_index_from_string(cb)
            return [self.cell(sheet, f"{get_column_letter(c)}{r}")
                    for r in range(int(ra), int(rb) + 1)
                    for c in range(c0, c1 + 1)]
        if CELL_RE.match(text):
            return self.cell(sheet, text)
        if text in self.names:
            return self.cell(*self.names[text])
        return XLError("#NAME?")

    def ev(self, n, sheet):
        kind = n[0]
        if kind == "num" or kind == "str":
            return n[1]
        if kind == "err":
            return XLError(n[1])
        if kind == "ref":
            return self.ref(n[1], sheet)
        if kind == "neg":
            return -_num(self.ev(n[1], sheet))
        if kind == "pct":
            return _num(self.ev(n[1], sheet)) / 100
        if kind == "bin":
            return self.binop(n[1], self.ev(n[2], sheet), self.ev(n[3], sheet))
        if kind == "func":
            return self.func(n[1], n[2], sheet)
        raise ValueError(kind)

    def binop(self, op, a, b):
        if op == "&":
            for v in (a, b):
                if isinstance(v, XLError):
                    raise _Propagate(v)
            return "".join(_fmt_concat(v) for v in (a, b))
        if op in ("=", "<>", "<", ">", "<=", ">="):
            if isinstance(a, str) or isinstance(b, str):
                x, y = str(a).lower(), str(b).lower()
            else:
                x, y = _num(a), _num(b)
            return {"=": x == y, "<>": x != y, "<": x < y, ">": x > y,
                    "<=": x <= y, ">=": x >= y}[op]
        x, y = _num(a), _num(b)
        if op == "+":
            return x + y
        if op == "-":
            return x - y
        if op == "*":
            return x * y
        if op == "/":
            if y == 0:
                raise _Propagate(XLError("#DIV/0!"))
            return x / y
        if op == "^":
            return x ** y
        raise ValueError(op)

    def func(self, name, args, sheet):
        ev = lambda i: self.ev(args[i], sheet)          # noqa: E731
        if name == "IF":
            cond = _num(ev(0))
            if cond:
                return ev(1)
            return ev(2) if len(args) > 2 else False
        if name == "IFERROR":
            try:
                v = ev(0)
                return ev(1) if isinstance(v, XLError) else v
            except _Propagate:
                return ev(1)
        if name == "CHOOSE":
            return ev(int(_num(ev(0))))
        vals = [self.ev(a, sheet) for a in args]
        if name in ("MAX", "MIN", "SUM", "AVERAGE"):
            flat = []
            for a, v in zip(args, vals):
                if isinstance(v, list):   # ranges skip blanks and text
                    flat += [_num(x) for x in _flatten(v)
                             if isinstance(x, (int, float)) or isinstance(x, XLError)]
                else:
                    flat.append(_num(v))
            if name == "SUM":
                return sum(flat)
            if name == "AVERAGE":
                return sum(flat) / len(flat)
            if not flat:
                return 0.0
            return max(flat) if name == "MAX" else min(flat)
        if name == "COUNT":
            return float(sum(1 for x in _flatten(vals)
                             if isinstance(x, (int, float)) and not isinstance(x, bool)))
        if name == "XIRR":
            return _xirr([_num(x) for x in _flatten([vals[0]])],
                         [_num(x) for x in _flatten([vals[1]])])
        if name == "AND":
            return all(_num(v) for v in _flatten(vals))
        if name == "OR":
            return any(_num(v) for v in _flatten(vals))
        if name == "NOT":
            return not _num(vals[0])
        if name == "ABS":
            return abs(_num(vals[0]))
        if name == "INT":
            return float(math.floor(_num(vals[0])))
        if name == "ROUND":
            return round(_num(vals[0]), int(_num(vals[1])))
        if name == "ROUNDUP":
            return _roundup(_num(vals[0]), _num(vals[1]))
        if name == "EDATE":
            return _edate(_num(vals[0]), _num(vals[1]))
        if name == "YEAR":
            return float(to_date(_num(vals[0])).year)
        if name == "MONTH":
            return float(to_date(_num(vals[0])).month)
        if name == "TEXT":
            return _text(vals[0], vals[1])
        if name == "INDEX":
            rng = vals[0] if isinstance(vals[0], list) else [vals[0]]
            return rng[int(_num(vals[1])) - 1]
        raise NotImplementedError(f"function {name}")


def _xirr(values, dates):
    """Annual rate r with sum(v / (1+r)^((d - d0)/365)) = 0, as Excel's XIRR.
    Bisection: slower than Excel's Newton step but cannot diverge."""
    def npv(r):
        return sum(v / (1 + r) ** ((d - dates[0]) / 365) for v, d in zip(values, dates))
    lo, hi = -0.9999, 100.0
    if npv(lo) * npv(hi) > 0:
        raise _Propagate(XLError("#NUM!"))
    for _ in range(300):
        mid = (lo + hi) / 2
        if npv(lo) * npv(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def _fmt_concat(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)
