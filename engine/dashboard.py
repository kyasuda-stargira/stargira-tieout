"""
Dashboard sheet: charts of the main figures over time, for every vertical.

The sheet holds a small linked table (one row per period, formulas only) and
line charts drawn from it. Flows (revenue, profit, cash flow) are shown as an
annual rate (period value × 12 / months), because periods change length from
monthly to quarterly to annual and raw period totals would jump at each change.
Balances (cash, net debt, ARR) are period-end values. The x-axis is a date axis,
so a quarter takes three times the width of a month.

Columns are found by their header text, so the same code serves the SaaS,
power and balance builders.
"""

from __future__ import annotations

from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.utils import get_column_letter

from build_model import BORDER, PCT1, _style_header

USD = "#,##0.00"
NUM1 = "#,##0.0"
# Categorical slots 1-3 (fixed order) from the shared chart palette.
SERIES_COLORS = ["2A78D6", "EB6834", "1BAF7A"]
TABLE_COL = 21          # data table starts in column U, right of the charts
CHART_W, CHART_H = 16, 7.5   # cm
ROWS_PER_CHART = 16


def _find(wb, sheet, header):
    ws = wb[sheet]
    for c in ws[1]:
        if c.value == header:
            return c.column_letter
    raise KeyError(f"{sheet}: no column headed {header!r}")


# Each vertical: list of (column header on Dashboard, number format, formula
# builder), then charts as (title, [Dashboard headers]). A formula builder gets
# (col, r) where col(sheet, header) returns the source column letter.
def _annual(sheet, header):
    return lambda col, r: f"{sheet}!{col(sheet, header)}{r}*12/Timeline!{col('Timeline', 'Months')}{r}"


def _end(sheet, header):
    return lambda col, r: f"{sheet}!{col(sheet, header)}{r}"


def _ratio(sheet, header):
    # KPI cells hold "" when undefined; NA() leaves a gap instead of a zero.
    return lambda col, r: (f'IF(ISNUMBER({sheet}!{col(sheet, header)}{r}), '
                           f'{sheet}!{col(sheet, header)}{r}, NA())')


def _fcf(col, r):
    return (f"(Cash_Flow!{col('Cash_Flow', 'CF from Ops ($M)')}{r}"
            f"+Cash_Flow!{col('Cash_Flow', 'CF from Inv ($M)')}{r})"
            f"*12/Timeline!{col('Timeline', 'Months')}{r}")


COMMON = [
    ("Revenue ($M/yr)", USD, _annual("PnL", "Revenue ($M)")),
    ("EBITDA ($M/yr)", USD, _annual("PnL", "EBITDA ($M)")),
    ("Net income ($M/yr)", USD, _annual("PnL", "Net Income ($M)")),
    ("Operating cash flow ($M/yr)", USD, _annual("Cash_Flow", "CF from Ops ($M)")),
    ("Free cash flow ($M/yr)", USD, _fcf),
    ("Cash ($M)", USD, _end("Balance_Sheet", "Cash ($M)")),
]
COMMON_CHARTS = [
    ("Revenue, EBITDA and net income ($M a year)",
     ["Revenue ($M/yr)", "EBITDA ($M/yr)", "Net income ($M/yr)"]),
    ("Operating and free cash flow ($M a year)",
     ["Operating cash flow ($M/yr)", "Free cash flow ($M/yr)"]),
    ("Cash and net debt ($M, period end)", ["Cash ($M)", "Net debt ($M)"]),
]


def _net_debt(debt_header):
    if debt_header is None:
        return lambda col, r: f"-Balance_Sheet!{col('Balance_Sheet', 'Cash ($M)')}{r}"
    return lambda col, r: (f"Balance_Sheet!{col('Balance_Sheet', debt_header)}{r}"
                           f"-Balance_Sheet!{col('Balance_Sheet', 'Cash ($M)')}{r}")


def _cum(sheet, header):
    def f(col, r, first):
        cur = f"{sheet}!{col(sheet, header)}{r}"
        return cur if r == first else f"{{prev}}+{cur}"
    f.cumulative = True
    return f


def _arr_per_head(col, r):
    heads = f"Headcount!{col('Headcount', 'Total Heads')}{r}"
    return f"IF({heads}=0, NA(), Revenue!{col('Revenue', 'ARR ($M)')}{r}*1000/{heads})"


VERTICALS = {
    "saas": ([("Net debt ($M)", USD, _net_debt(None)),
              ("ARR ($M)", USD, _end("Revenue", "ARR ($M)")),
              ("Customers", "#,##0", _end("Customers", "Customers Close")),
              ("Headcount", "#,##0", _end("Headcount", "Total Heads")),
              ("ARR per head ($k)", "#,##0", _arr_per_head),
              ("Unplanned funding, cumulative ($M)", USD, _cum("Funding", "Unplanned Funding ($M)")),
              ("NRR (annualised)", PCT1, _ratio("KPI", "NRR (annualised)")),
              ("Gross margin", PCT1, _ratio("KPI", "Gross Margin")),
              ("CAC payback (months)", NUM1, _ratio("KPI", "CAC Payback (months)"))],
             [("ARR ($M, period end)", ["ARR ($M)"]),
              ("Customers and headcount (period end)", ["Customers", "Headcount"]),
              ("ARR per head ($k)", ["ARR per head ($k)"]),
              ("Unplanned funding, cumulative ($M)", ["Unplanned funding, cumulative ($M)"]),
              ("NRR and gross margin", ["NRR (annualised)", "Gross margin"]),
              ("CAC payback (months)", ["CAC payback (months)"])]),
    "power": ([("Net debt ($M)", USD, _net_debt("Debt ($M)")),
               ("DSCR", "0.00", _ratio("Debt", "DSCR"))],
              [("DSCR (cash available ÷ debt service)", ["DSCR"])]),
    "balance": ([("Net debt ($M)", USD, _net_debt(None)),
                 ("AUM ($M)", USD, _end("KPI", "AUM ($M)")),
                 ("Runway on net burn (months)", NUM1, _ratio("KPI", "Runway on Net Burn (mo)"))],
                [("AUM ($M, period end)", ["AUM ($M)"]),
                 ("Runway on net burn (months)", ["Runway on net burn (months)"])]),
}


def build(wb, kind):
    extra_cols, extra_charts = VERTICALS[kind]
    cols = COMMON + extra_cols
    charts = COMMON_CHARTS + extra_charts
    last = wb["Timeline"].max_row
    first = 2

    def col(sheet, header):
        return _find(wb, sheet, header)

    ws = wb.create_sheet("Dashboard", 0)
    start_col = col("Timeline", "Period Start")
    label_col = col("Timeline", "Period Label")
    heads = [("Period start", "mmm yyyy", lambda c, r: f"Timeline!{start_col}{r}"),
             ("Period", None, lambda c, r: f"Timeline!{label_col}{r}")] + cols
    where = {}
    for i, (h, fmt, fn) in enumerate(heads):
        ci = TABLE_COL + i
        letter = get_column_letter(ci)
        where[h] = ci
        ws.column_dimensions[letter].width = max(12, len(h) + 2)
        _style_header(ws.cell(row=1, column=ci, value=h))
        for r in range(first, last + 1):
            if getattr(fn, "cumulative", False):
                f = fn(col, r, first).replace("{prev}", f"{letter}{r-1}")
            else:
                f = fn(col, r)
            c = ws.cell(row=r, column=ci, value="=" + f)
            c.border = BORDER
            if fmt:
                c.number_format = fmt

    dates = Reference(ws, min_col=where["Period start"], min_row=first, max_row=last)
    for n, (title, series) in enumerate(charts):
        ch = LineChart()
        ch.title = title
        ch.width, ch.height = CHART_W, CHART_H
        ch.legend = ch.legend if len(series) > 1 else None
        if ch.legend is not None:
            ch.legend.position = "b"
        for k, h in enumerate(series):
            ch.add_data(Reference(ws, min_col=where[h], min_row=1, max_row=last),
                        titles_from_data=True)
            s = ch.series[-1]
            s.smooth = False
            s.graphicalProperties.line.solidFill = SERIES_COLORS[k]
            s.graphicalProperties.line.width = 25400   # 2 pt
        ch.set_categories(dates)
        ch.x_axis = DateAxis(crossAx=100)
        ch.x_axis.number_format = "yyyy"
        ch.x_axis.majorTimeUnit = "years"
        if kind == "power":   # 25+ years: label every fifth
            ch.x_axis.majorUnit = 5
        ch.x_axis.tickLblPos = "low"   # keep year labels below negative values
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.crossAx = 500
        fmt = next(f for h2, f, _ in heads if h2 == series[0])
        if fmt:
            ch.y_axis.number_format = fmt
        anchor_col = "A" if n % 2 == 0 else "K"
        ws.add_chart(ch, f"{anchor_col}{1 + (n // 2) * ROWS_PER_CHART}")
    ws.sheet_view.zoomScale = 90
    wb.active = 0
