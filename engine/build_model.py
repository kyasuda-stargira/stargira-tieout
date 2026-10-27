#!/usr/bin/env python3
"""
Financial Model Generator — full 3-statement model.

Reads:  <schema.yaml> + <answers.yaml>
Emits:  [output.xlsx], default output/<answers name>.xlsx

The schema's meta.archetypes picks the revenue logic: subscription ->
vertical_saas.py, capacity -> vertical_power.py, otherwise the balance model
built in this file (12 sheets, template verticals/balance.yaml).

Balance model: revenue is a yield on a balance the business holds (for example a
lending book or assets under management), plus transaction fees and other revenue.
Main simplifications:
  - The balance follows straight lines between three milestones; there is no
    customer-level or cohort detail.
  - Perfect-foresight cash policy: each period closes with enough cash to cover
    the next period's operating shortfall, using that period's known cash flow.
    Equity is injected only when needed and dividends are paid only above the
    target. A real business would forecast this with error, not know it.
  - Injected and retained cash earns nothing; yield accrues on the balance only.
  - No credit losses, provisions or funding cost: the yield is treated as net
    revenue. For a lender, model these outside or treat the yield as net of them.
  - Own-funds buffer = MAX(fixed floor, % of end-of-period balance). It is one
    capital requirement, funded by equity and held as a separate asset; there is
    no tiering of the assets behind it.
  - OpEx moves in a straight line in time (months from launch) from its launch
    level to its final level, reached in the last period.
  - Tax losses carry forward without limit.
  - A balance milestone after the last period is never reached; the build
    warns and the KPI sheet's Timeline Check cell says so.

Architectural invariants:
  1. Single timeline anchor                  (launch_anchor_date named range)
  2. Single source per assumption            (every input -> one named range)
  3. Zero hardcoded numbers outside Assumptions
  4. Single cash policy                      (one min_cash floor)
  5. Perfect-foresight cash policy           (fund/distribute vs next-period oper CF)
  6. Balance Sheet as ground truth           (A = L + E proved per period)
  7. Cash Flow reconciles to BS              (CF closing = BS cash proved per period)
"""

from __future__ import annotations
import sys
from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

# ---------------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------------
HEADER_FILL  = PatternFill("solid", fgColor="1F3864")
HEADER_FONT  = Font(bold=True, color="FFFFFF")
INPUT_FILL   = PatternFill("solid", fgColor="FFF2CC")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
CHECK_GOOD   = PatternFill("solid", fgColor="E2EFDA")
THIN   = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

UNIT_FMT = {
    "usd_millions": "#,##0.00",
    "integer":      "0",
    "date":         "mmm yyyy",
    "percent":      "0.00%",
    "ratio":        "0.00",
    "bps":          "0",
    "text":         "@",
}

USD3 = "#,##0.000"
USD2 = "#,##0.00"
PCT1 = "0.0%"


def _style_header(cell):
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER


def _c(ws, row, col, val, fmt=USD3):
    """Write a cell with border + number format."""
    c = ws.cell(row=row, column=col, value=val)
    c.number_format = fmt
    c.border = BORDER
    return c


def _headers(ws, cols):
    """Write header row and set column widths."""
    for i, (_, label, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        _style_header(ws.cell(row=1, column=i, value=label))


def _link(ws, row):
    """Link col A (period_id) and col B (period_label) to Timeline."""
    ws.cell(row=row, column=1, value=f"=Timeline!A{row}").border = BORDER
    ws.cell(row=row, column=2, value=f"=Timeline!C{row}").border = BORDER


def _months(ws, row):
    """Link col C (months) to Timeline."""
    ws.cell(row=row, column=3, value=f"=Timeline!F{row}").border = BORDER


# ---------------------------------------------------------------------------
# Column definitions — one tuple per sheet
# ---------------------------------------------------------------------------
TIMELINE_COLS = [
    ("period_id",          "Period ID",            12),
    ("period_date",        "Period Start",         14),
    ("period_label",       "Period Label",         14),
    ("period_year",        "Year",                  8),
    ("period_type",        "Type",                 12),
    ("period_months",      "Months",                8),
    ("months_from_launch", "Mths from Launch",     16),
]

CIRC_COLS = [
    ("period_id",   "Period ID",        12),
    ("period_label","Period Label",     14),
    ("months",      "Months",           8),
    ("circ_start",  "AUM Start ($M)",  16),
    ("circ_end",    "AUM End ($M)",    16),
    ("circ_avg",    "AUM Avg ($M)",    16),
]

REV_COLS = [
    ("period_id",   "Period ID",         12),
    ("period_label","Period Label",      14),
    ("months",      "Months",            8),
    ("yield_rev",   "Yield Rev ($M)",   16),
    ("fee_rev",     "Fee Rev ($M)",     14),
    ("other_rev",   "Other Rev ($M)",   16),
    ("total_rev",   "Total Rev ($M)",   16),
]

OPEX_COLS = [
    ("period_id",    "Period ID",         12),
    ("period_label", "Period Label",      14),
    ("months",       "Months",            8),
    ("monthly_burn", "Monthly Burn ($M)", 18),
    ("opex_period",  "OpEx ($M)",         14),
]

WC_COLS = [
    ("period_id",   "Period ID",          12),
    ("period_label","Period Label",       14),
    ("months",      "Months",             8),
    ("ar",          "Accts Recv ($M)",   16),
    ("prepaid",     "Prepaid ($M)",      14),
    ("ap",          "Accts Pay ($M)",    16),
    ("accrued",     "Accrued Exp ($M)",  16),
    ("net_wc",      "Net WC ($M)",       14),
    ("delta_wc",    "ΔWC ($M)",     12),
]

PNL_COLS = [
    ("period_id",   "Period ID",           12),
    ("period_label","Period Label",        14),
    ("revenue",     "Revenue ($M)",        16),
    ("opex",        "OpEx ($M)",           14),
    ("ebitda",      "EBITDA ($M)",         14),
    ("pre_tax",     "Pre-tax Inc ($M)",    18),
    ("nol_open",    "Tax Losses Open ($M)", 20),
    ("nol_used",    "Tax Losses Used ($M)", 20),
    ("nol_close",   "Tax Losses Close ($M)", 20),
    ("tax",         "Tax ($M)",            12),
    ("net_income",  "Net Income ($M)",     16),
]

FC_COLS = [
    ("period_id",         "Period ID",             12),
    ("period_label",      "Period Label",          14),
    ("cash_opening",      "Cash Open ($M)",        16),
    ("net_income",        "Net Income ($M)",       16),
    ("delta_wc",          "ΔWC ($M)",         12),
    ("operating_cf",      "Oper CF ($M)",          14),
    ("own_funds_buffer",  "Own-Funds Buffer ($M)", 20),
    ("delta_buffer",      "ΔBuffer ($M)",     14),
    ("cash_before",       "Pre-Fund ($M)",         16),
    ("next_oper_cf",      "Next Oper CF ($M)",     16),
    ("target_cash",       "Target Cash ($M)",      16),
    ("capital_injection", "Injection ($M)",        16),
    ("dividend",          "Dividend ($M)",         14),
    ("cash_closing",      "Cash Close ($M)",       16),
]

BS_COLS = [
    ("period_id",       "Period ID",            12),
    ("period_label",    "Period Label",         14),
    ("cash",            "Cash ($M)",            14),
    ("own_funds_buf",   "Own-Funds Buf ($M)",   18),
    ("ar",              "Accts Recv ($M)",      16),
    ("prepaid",         "Prepaid ($M)",         14),
    ("total_assets",    "Total Assets ($M)",    18),
    ("ap",              "Accts Pay ($M)",       16),
    ("accrued",         "Accrued Exp ($M)",     16),
    ("total_liab",      "Total Liab ($M)",      16),
    ("paid_in_capital", "Paid-in Cap ($M)",     16),
    ("retained_earn",   "Retained Earn ($M)",   18),
    ("total_equity",    "Total Equity ($M)",    18),
    ("total_le",        "Total L+E ($M)",       16),
    ("balance_check",   "BS Check ($M)",        14),
]

CF_COLS = [
    ("period_id",     "Period ID",            12),
    ("period_label",  "Period Label",         14),
    ("net_income",    "Net Income ($M)",      16),
    ("delta_ar",      "ΔAR ($M)",        12),
    ("delta_prepaid", "ΔPrepaid ($M)",   14),
    ("delta_ap",      "ΔAP ($M)",        12),
    ("delta_accrued", "ΔAccrued ($M)",   14),
    ("cfo",           "CF from Ops ($M)",     16),
    ("cfi",           "CF from Inv ($M)",     16),
    ("cff",           "CF from Fin ($M)",     16),
    ("net_change",    "Net ΔCash ($M)",  16),
    ("opening_cash",  "Open Cash ($M)",       16),
    ("closing_cash",  "Close Cash ($M)",      16),
    ("cf_check",      "CF Check ($M)",        14),
]

KPI_COLS = [
    ("period_id",     "Period ID",            12),
    ("period_label",  "Period Label",         14),
    ("aum",           "AUM ($M)",             14),
    ("revenue",       "Revenue ($M)",         16),
    ("opex",          "OpEx ($M)",            14),
    ("ebitda_margin", "EBITDA Margin",        14),
    ("net_income",    "Net Income ($M)",      16),
    ("cum_injection", "Cum Inject ($M)",      16),
    ("cash",          "Cash ($M)",            14),
    ("runway",        "Runway on Net Burn (mo)", 22),
]


# ═══════════════════════════════════════════════════════════════════════════
# Schema
# ═══════════════════════════════════════════════════════════════════════════

def schema_assumptions(schema: dict) -> list[dict]:
    """Every input the schema defines, one dict per named range.
    verify_model.py audits named ranges against this list.

    Vertical templates may also carry two table blocks, flattened here into
    ordinary inputs so they land on the Assumptions sheet like any other:
      scenarios: one input per (axis, scenario), id scn_<axis>_<scenario id>
      curves:    one input per (curve, point, field), id crv_<curve>_<n>_<field>
    """
    items = list(schema["assumptions"])
    for axis in scenario_axes(schema):
        for s in scenario_rows(schema):
            items.append({"id": f"scn_{axis}_{s['id'].lower()}",
                          "label": f"{s['id']} scenario: {axis}",
                          "section": "Scenarios", "unit": "ratio",
                          "default": s[axis]})
    for name, points in (schema.get("curves") or {}).items():
        for n, p in enumerate(points, 1):
            for field, v in p.items():
                items.append({"id": f"crv_{name}_{n}_{field}",
                              "label": f"{name}, point {n}: {field}",
                              "section": "Curves",
                              "unit": "integer" if field == "year" else "ratio",
                              "default": v})
    return items


def scenario_rows(schema: dict) -> list[dict]:
    """Scenario rows ordered by their switch code (1, 2, 3...)."""
    return sorted(schema.get("scenarios") or [], key=lambda s: s["code"])


def scenario_axes(schema: dict) -> list[str]:
    rows = scenario_rows(schema)
    return [k for k in rows[0] if k not in ("id", "code")] if rows else []


# ═══════════════════════════════════════════════════════════════════════════
# Sheet builders
# ═══════════════════════════════════════════════════════════════════════════

# ---------------------------------------------------------------------------
# 1. Assumptions — every input lives here as a named range
# ---------------------------------------------------------------------------
def build_assumptions(wb: Workbook, schema: dict, answers: dict) -> dict[str, str]:
    ws = wb.create_sheet("Assumptions")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 56
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 16

    for col, hdr in enumerate(["#", "Assumption", "ID", "Unit", "Value"], 1):
        _style_header(ws.cell(row=1, column=col, value=hdr))

    refs: dict[str, str] = {}
    row = 2
    last_section = None
    n = 0

    for a in schema_assumptions(schema):
        if a["section"] != last_section:
            c = ws.cell(row=row, column=2, value=a["section"])
            c.font = Font(bold=True)
            c.fill = SECTION_FILL
            for col in range(1, 6):
                ws.cell(row=row, column=col).fill = SECTION_FILL
            row += 1
            last_section = a["section"]

        n += 1
        ws.cell(row=row, column=1, value=n).border = BORDER
        ws.cell(row=row, column=2, value=a["label"]).border = BORDER
        ws.cell(row=row, column=3, value=a["id"]).border = BORDER
        ws.cell(row=row, column=4, value=a["unit"]).border = BORDER

        value = answers.get(a["id"], a.get("default"))
        if a["unit"] == "date" and isinstance(value, str):
            value = datetime.fromisoformat(value).date()

        vcell = ws.cell(row=row, column=5, value=value)
        vcell.fill = INPUT_FILL
        vcell.border = BORDER
        fmt = UNIT_FMT.get(a["unit"])
        if fmt:
            vcell.number_format = fmt

        ref = f"Assumptions!${get_column_letter(5)}${row}"
        refs[a["id"]] = ref
        wb.defined_names[a["id"]] = DefinedName(name=a["id"], attr_text=ref)
        row += 1

    ws.freeze_panes = "A2"
    return refs


# ---------------------------------------------------------------------------
# 2. Timeline — every cell a formula keyed off launch_anchor_date
# ---------------------------------------------------------------------------
def build_timeline(wb: Workbook, schema: dict, answers: dict) -> dict:
    ws = wb.create_sheet("Timeline")
    _headers(ws, TIMELINE_COLS)

    n_periods = answers["monthly_periods"] + answers["quarterly_periods"]
    monthly_n = answers["monthly_periods"]
    pre_n     = answers["pre_launch_periods"]

    for i in range(n_periods):
        row = i + 2
        pid = i - pre_n

        # A: period_id
        ws.cell(row=row, column=1, value=pid).border = BORDER

        # B: period_date via EDATE
        if i < monthly_n:
            date_f = f"=EDATE(launch_anchor_date, {pid})"
        else:
            q = i - monthly_n
            date_f = f"=EDATE(launch_anchor_date, (monthly_periods - pre_launch_periods) + {q}*3)"
        c = ws.cell(row=row, column=2, value=date_f)
        c.number_format = "mmm yyyy"
        c.border = BORDER

        # C: period_label
        if i < monthly_n:
            lbl = f'=TEXT(B{row}, "mmm yyyy")'
        else:
            # "Q4 2027" only when the period is a calendar quarter; otherwise
            # name its months, e.g. "Nov 27–Jan 28".
            lbl = (f'=IF(MONTH(B{row})=ROUNDUP(MONTH(B{row})/3,0)*3-2, '
                   f'"Q" & ROUNDUP(MONTH(B{row})/3,0) & " " & YEAR(B{row}), '
                   f'TEXT(B{row}, "mmm yy") & "–" & TEXT(EDATE(B{row}, 2), "mmm yy"))')
        ws.cell(row=row, column=3, value=lbl).border = BORDER

        # D: year
        ws.cell(row=row, column=4, value=f"=YEAR(B{row})").border = BORDER

        # E: type
        if pid < 0:
            t = '="pre_launch"'
        elif i < monthly_n:
            t = '="monthly"'
        else:
            t = '="quarterly"'
        ws.cell(row=row, column=5, value=t).border = BORDER

        # F: period_months
        ws.cell(row=row, column=6, value=1 if i < monthly_n else 3).border = BORDER

        # G: months_from_launch (computed from dates)
        mfl = (f"=(YEAR(B{row})-YEAR(launch_anchor_date))*12"
               f"+MONTH(B{row})-MONTH(launch_anchor_date)")
        ws.cell(row=row, column=7, value=mfl).border = BORDER

    ws.freeze_panes = "A2"
    return {"first": 2, "last": n_periods + 1, "n": n_periods}


# ---------------------------------------------------------------------------
# 3. Balance ramp — AUM curve via piecewise-linear interpolation.
#    Milestones (year + value) are all inputs; knee points = milestone_year*12.
#    Pre-launch = 0.  (The only literals are calendar 12 and the 0 guard.)
# ---------------------------------------------------------------------------
def build_balance_ramp(wb: Workbook, ri: dict) -> None:
    ws = wb.create_sheet("Balance_Ramp")
    _headers(ws, CIRC_COLS)

    # milestone breakpoints in months, all derived from named ranges
    b1 = "balance_milestone1_year*12"
    b2 = "balance_milestone2_year*12"
    b3 = "balance_milestone3_year*12"
    v0 = "balance_launch"
    v1 = "balance_milestone1_value"
    v2 = "balance_milestone2_value"
    v3 = "balance_milestone3_value"

    first, last = ri["first"], ri["last"]
    for row in range(first, last + 1):
        _link(ws, row)
        _months(ws, row)

        # D: circ_start
        if row == first:
            _c(ws, row, 4, f"=IF(A{row}<0, 0, {v0})")
        else:
            _c(ws, row, 4,
               f"=IF(A{row}<0, 0, IF(A{row}=0, {v0}, E{row-1}))")

        # E: circ_end — piecewise-linear through the 3 milestones
        m = f"(Timeline!G{row}+Timeline!F{row})"
        _c(ws, row, 5,
           f"=IF({m}<=0, 0, "
           f"IF({m}<={b1}, "
           f"{v0}+({v1}-{v0})*{m}/({b1}), "
           f"IF({m}<={b2}, "
           f"{v1}+({v2}-{v1})*({m}-{b1})/(({b2})-({b1})), "
           f"{v2}+({v3}-{v2})*MIN({m}-{b2},({b3})-({b2}))/(({b3})-({b2})))))")

        # F: circ_avg
        _c(ws, row, 6, f"=(D{row}+E{row})/2")

    ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------
# 4. Revenue — yield on the balance + transaction fees + other
# ---------------------------------------------------------------------------
def build_revenue(wb: Workbook, ri: dict) -> None:
    ws = wb.create_sheet("Revenue")
    _headers(ws, REV_COLS)

    for row in range(ri["first"], ri["last"] + 1):
        _link(ws, row)
        _months(ws, row)

        # D: yield_revenue = avg_circ * rate * months / 12
        _c(ws, row, 4, f"=Balance_Ramp!F{row}*asset_yield_rate*C{row}/12")

        # E: fee_revenue = avg_circ * velocity * months/12 * bps/10000
        _c(ws, row, 5,
           f"=Balance_Ramp!F{row}*transaction_velocity*C{row}/12"
           f"*transaction_fee_bps/10000")

        # F: other_revenue
        _c(ws, row, 6, f"=other_revenue_annual*C{row}/12")

        # G: total
        _c(ws, row, 7, f"=D{row}+E{row}+F{row}")

    ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------
# 5. OpEx — per-period burn, linear ramp from launch to terminal
# ---------------------------------------------------------------------------
def build_opex(wb: Workbook, ri: dict) -> None:
    ws = wb.create_sheet("OpEx")
    _headers(ws, OPEX_COLS)

    first, last = ri["first"], ri["last"]
    # Straight line in time: months from launch at the start of this period,
    # over months from launch at the start of the last period. (Period ID
    # would make a three-month quarter count as one step.)
    denom = f"MAX(1, Timeline!$G${last})"

    for row in range(first, last + 1):
        _link(ws, row)
        _months(ws, row)

        # D: monthly_burn
        _c(ws, row, 4,
           f"=IF(A{row}<0, opex_prelaunch_monthly, "
           f"opex_initial_monthly + (opex_terminal_monthly - opex_initial_monthly)"
           f" * Timeline!G{row} / {denom})")

        # E: opex_period = burn * months
        _c(ws, row, 5, f"=D{row}*C{row}")

    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# 6. Working Capital — AR, AP, prepaid, accrued
# ---------------------------------------------------------------------------
def build_working_capital(wb: Workbook, ri: dict) -> None:
    ws = wb.create_sheet("Working_Capital")
    _headers(ws, WC_COLS)

    first = ri["first"]
    for row in range(first, ri["last"] + 1):
        _link(ws, row)
        _months(ws, row)

        # D: AR = monthly_revenue * ar_days / 30
        _c(ws, row, 4, f"=IF(C{row}=0, 0, Revenue!G{row}/C{row}*ar_days/30)")

        # E: Prepaid = monthly_burn * prepaid_months
        _c(ws, row, 5, f"=OpEx!D{row}*prepaid_months")

        # F: AP = monthly_burn * ap_days / 30
        _c(ws, row, 6, f"=OpEx!D{row}*ap_days/30")

        # G: Accrued = monthly_burn * accrued_expense_pct
        _c(ws, row, 7, f"=OpEx!D{row}*accrued_expense_pct")

        # H: Net WC = (AR + Prepaid) - (AP + Accrued)
        _c(ws, row, 8, f"=(D{row}+E{row})-(F{row}+G{row})")

        # I: Delta WC
        if row == first:
            _c(ws, row, 9, f"=H{row}")
        else:
            _c(ws, row, 9, f"=H{row}-H{row-1}")

    ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------
# 7. P&L — income statement
# ---------------------------------------------------------------------------
def build_pnl(wb: Workbook, ri: dict) -> None:
    ws = wb.create_sheet("PnL")
    _headers(ws, PNL_COLS)

    first = ri["first"]
    for row in range(first, ri["last"] + 1):
        _link(ws, row)

        _c(ws, row, 3, f"=Revenue!G{row}", USD2)                         # revenue
        _c(ws, row, 4, f"=OpEx!E{row}", USD2)                            # opex
        _c(ws, row, 5, f"=C{row}-D{row}", USD2)                          # ebitda
        _c(ws, row, 6, f"=E{row}", USD2)                                 # pre_tax
        # Tax losses carried forward, as in the SaaS and power models:
        # losses build up while pre-tax is negative and are used, without
        # limit, against later profits before any tax is charged.
        if row == first:
            _c(ws, row, 7, 0, USD2)                                      # nol_open
        else:
            _c(ws, row, 7, f"=I{row-1}", USD2)                           # nol_open
        _c(ws, row, 8, f"=IF(F{row}>0, MIN(G{row}, F{row}), 0)", USD2)   # nol_used
        _c(ws, row, 9, f"=G{row}-H{row}+MAX(0, -F{row})", USD2)          # nol_close
        _c(ws, row, 10, f"=MAX(0, F{row}-H{row})*corporate_tax_rate", USD2)  # tax
        _c(ws, row, 11, f"=F{row}-J{row}", USD2)                         # net_income

    ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------
# 8. Forward Cash Requirements — perfect-foresight cash policy.
#    The simulation knows next period's operating CF, so the funding decision
#    is non-circular: injected/retained cash earns nothing (revenue accrues
#    on the balance), and each period closes at a target that pre-funds the
#    next period's operating shortfall:
#      target    = MAX(floor, floor - next_oper_CF)
#      injection = MAX(0, target - cash)            (next oper CF negative)
#      dividend  = MIN(payout_ratio x NI, cash - target, retained earnings)
#      C cash_open | D NI | E ΔWC | F oper_CF | G buffer_req | H Δbuffer
#      I cash_before | J next_oper_cf | K target | L injection | M dividend
#      N cash_close
# ---------------------------------------------------------------------------
def build_forward_cash(wb: Workbook, ri: dict) -> None:
    ws = wb.create_sheet("Forward_Cash")
    _headers(ws, FC_COLS)

    first, last = ri["first"], ri["last"]
    for row in range(first, last + 1):
        _link(ws, row)

        # C: cash_opening
        if row == first:
            _c(ws, row, 3, "=initial_seed_deposit", USD2)
        else:
            _c(ws, row, 3, f"=N{row-1}", USD2)

        # D: net_income
        _c(ws, row, 4, f"=PnL!K{row}", USD2)

        # E: delta_wc
        _c(ws, row, 5, f"=Working_Capital!I{row}", USD2)

        # F: operating_cf = NI - delta_WC
        _c(ws, row, 6, f"=D{row}-E{row}", USD2)

        # G: own_funds_buffer = MAX(floor, pct x end-of-period AUM)
        _c(ws, row, 7,
           f"=MAX(own_funds_floor, own_funds_pct*Balance_Ramp!E{row})", USD2)

        # H: delta_buffer (cash moved into the regulatory buffer this period)
        if row == first:
            _c(ws, row, 8, f"=G{row}", USD2)
        else:
            _c(ws, row, 8, f"=G{row}-G{row-1}", USD2)

        # I: cash_before_injection = opening + operating_CF - buffer funding
        _c(ws, row, 9, f"=C{row}+F{row}-H{row}", USD2)

        # J: next period's operating CF (last period: assume own repeats)
        nxt = row + 1 if row < last else row
        _c(ws, row, 10, f"=F{nxt}", USD2)

        # K: target closing cash — pre-fund next period's operating
        #    shortfall, never below the minimum floor
        _c(ws, row, 11,
           f"=MAX(minimum_operating_cash, minimum_operating_cash-J{row})", USD2)

        # L: capital_injection = MAX(0, target - cash_before)
        _c(ws, row, 12, f"=MAX(0, K{row}-I{row})", USD2)

        # M: dividend = payout_ratio x NI, capped by excess cash above target
        #    and by pre-dividend retained earnings (no distributions until
        #    cumulatively profitable)
        re_pre = f"D{row}" if row == first else f"Balance_Sheet!L{row-1}+D{row}"
        _c(ws, row, 13,
           f"=MIN(dividend_payout_ratio*MAX(0, D{row}), "
           f"MAX(0, I{row}-K{row}), MAX(0, {re_pre}))", USD2)

        # N: cash_closing = cash_before + injection - dividend
        _c(ws, row, 14, f"=I{row}+L{row}-M{row}", USD2)

    ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------
# 9. Balance Sheet — Assets = Liabilities + Equity
# ---------------------------------------------------------------------------
def build_balance_sheet(wb: Workbook, ri: dict) -> None:
    ws = wb.create_sheet("Balance_Sheet")
    _headers(ws, BS_COLS)

    first = ri["first"]
    for row in range(first, ri["last"] + 1):
        _link(ws, row)

        # Assets   C cash | D own-funds buffer | E AR | F prepaid | G total
        _c(ws, row, 3,  f"=Forward_Cash!N{row}", USD2)             # operating cash
        _c(ws, row, 4,  f"=Forward_Cash!G{row}", USD2)             # own-funds buffer
        _c(ws, row, 5,  f"=Working_Capital!D{row}", USD2)          # AR
        _c(ws, row, 6,  f"=Working_Capital!E{row}", USD2)          # prepaid
        _c(ws, row, 7,  f"=C{row}+D{row}+E{row}+F{row}", USD2)    # total_assets

        # Liabilities   H AP | I accrued | J total_liab
        _c(ws, row, 8,  f"=Working_Capital!F{row}", USD2)          # AP
        _c(ws, row, 9,  f"=Working_Capital!G{row}", USD2)          # accrued
        _c(ws, row, 10, f"=H{row}+I{row}", USD2)                  # total_liab

        # Equity — K paid-in capital (cumulative injections incl. seed)
        if row == first:
            _c(ws, row, 11, f"=initial_seed_deposit+Forward_Cash!L{row}", USD2)
        else:
            _c(ws, row, 11, f"=K{row-1}+Forward_Cash!L{row}", USD2)

        # Equity — L retained earnings (cumulative NI less dividends)
        if row == first:
            _c(ws, row, 12, f"=PnL!K{row}-Forward_Cash!M{row}", USD2)
        else:
            _c(ws, row, 12, f"=L{row-1}+PnL!K{row}-Forward_Cash!M{row}", USD2)

        _c(ws, row, 13, f"=K{row}+L{row}", USD2)                  # total_equity
        _c(ws, row, 14, f"=J{row}+M{row}", USD2)                  # total L+E

        # BS check (should be 0.00):  total_assets - total_L+E
        chk = _c(ws, row, 15, f"=G{row}-N{row}", USD2)
        chk.fill = CHECK_GOOD

    ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------
# 10. Cash Flow — indirect method
# ---------------------------------------------------------------------------
def build_cash_flow(wb: Workbook, ri: dict) -> None:
    ws = wb.create_sheet("Cash_Flow")
    _headers(ws, CF_COLS)

    first = ri["first"]
    for row in range(first, ri["last"] + 1):
        _link(ws, row)

        # C: net_income
        _c(ws, row, 3, f"=PnL!K{row}", USD2)

        # D-G: working-capital adjustments
        if row == first:
            _c(ws, row, 4, f"=-Working_Capital!D{row}", USD2)
            _c(ws, row, 5, f"=-Working_Capital!E{row}", USD2)
            _c(ws, row, 6, f"=Working_Capital!F{row}", USD2)
            _c(ws, row, 7, f"=Working_Capital!G{row}", USD2)
        else:
            _c(ws, row, 4, f"=-(Working_Capital!D{row}-Working_Capital!D{row-1})", USD2)
            _c(ws, row, 5, f"=-(Working_Capital!E{row}-Working_Capital!E{row-1})", USD2)
            _c(ws, row, 6, f"=Working_Capital!F{row}-Working_Capital!F{row-1}", USD2)
            _c(ws, row, 7, f"=Working_Capital!G{row}-Working_Capital!G{row-1}", USD2)

        # H: CFO
        _c(ws, row, 8, f"=C{row}+D{row}+E{row}+F{row}+G{row}", USD2)

        # I: CFI (investing) — cash moved into the own-funds buffer
        _c(ws, row, 9, f"=-Forward_Cash!H{row}", USD2)

        # J: CFF (capital injections less dividends)
        if row == first:
            _c(ws, row, 10,
               f"=initial_seed_deposit+Forward_Cash!L{row}-Forward_Cash!M{row}", USD2)
        else:
            _c(ws, row, 10, f"=Forward_Cash!L{row}-Forward_Cash!M{row}", USD2)

        # K: net change in cash
        _c(ws, row, 11, f"=H{row}+I{row}+J{row}", USD2)

        # L: opening cash
        if row == first:
            _c(ws, row, 12, 0, USD2)
        else:
            _c(ws, row, 12, f"=M{row-1}", USD2)

        # M: closing cash
        _c(ws, row, 13, f"=L{row}+K{row}", USD2)

        # N: CF check vs BS operating cash (should be 0.00)
        chk = _c(ws, row, 14, f"=M{row}-Balance_Sheet!C{row}", USD2)
        chk.fill = CHECK_GOOD

    ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------
# 11. KPI Dashboard
# ---------------------------------------------------------------------------
def build_kpi(wb: Workbook, ri: dict) -> None:
    ws = wb.create_sheet("KPI")
    _headers(ws, KPI_COLS)

    for row in range(ri["first"], ri["last"] + 1):
        _link(ws, row)

        _c(ws, row, 3,  f"=Balance_Ramp!E{row}", USD2)                  # AUM
        _c(ws, row, 4,  f"=PnL!C{row}", USD2)                          # revenue
        _c(ws, row, 5,  f"=PnL!D{row}", USD2)                          # opex
        _c(ws, row, 6,  f'=IF(D{row}=0, "n/a", (D{row}-E{row})/D{row})', PCT1)  # margin
        _c(ws, row, 7,  f"=PnL!K{row}", USD2)                          # net income
        _c(ws, row, 8,  f"=Balance_Sheet!K{row}", USD2)                 # cum inject
        _c(ws, row, 9,  f"=Balance_Sheet!C{row}", USD2)                 # cash
        # Runway on net burn: cash / monthly operating cash outflow;
        # "n/a" when operating cash flow is not negative (nothing burned).
        _c(ws, row, 10, f'=IF(Forward_Cash!F{row}>=0, "n/a", '
                        f'I{row}*Timeline!F{row}/(-Forward_Cash!F{row}))', "0.0")

    # One timeline check: a balance milestone after the last period is never
    # reached, so the balance ends below its target.
    last = ri["last"]
    _style_header(ws.cell(row=1, column=12, value="Timeline Check"))
    ws.column_dimensions["L"].width = 60
    ws.cell(row=2, column=12, value=(
        f'=IF(MAX(balance_milestone1_year, balance_milestone2_year, '
        f'balance_milestone3_year)*12>Timeline!G{last}+Timeline!F{last}, '
        f'"WARNING: a balance milestone falls after the last period, so the '
        f'balance never reaches it. Add periods or move the milestone.", '
        f'"OK: every balance milestone falls inside the timeline")')).border = BORDER

    ws.freeze_panes = "C2"


def warn_milestones(values: dict) -> None:
    """Build-time twin of the KPI Timeline Check."""
    end = (values["monthly_periods"] - values["pre_launch_periods"]
           + 3 * values["quarterly_periods"])
    for k in (1, 2, 3):
        yr = values[f"balance_milestone{k}_year"]
        if yr * 12 > end:
            print(f"WARNING: balance milestone {k} (year {yr}, month {yr * 12}) falls "
                  f"after the last period (month {end} from launch); the balance "
                  f"will not reach balance_milestone{k}_value.", file=sys.stderr)


# ---------------------------------------------------------------------------
# 12. Manifest (hidden)
# ---------------------------------------------------------------------------
def build_manifest(wb: Workbook, schema: dict, answers_path: Path) -> None:
    ws = wb.create_sheet("_Manifest")
    ws["A1"] = "Generator"
    ws["B1"] = "financial_model_generator v0.2 — full 3-statement"
    ws["A2"] = "Schema version"
    ws["B2"] = schema["meta"]["schema_version"]
    ws["A3"] = "Vertical"
    ws["B3"] = schema["meta"].get("vertical", "")
    ws["A4"] = "Slice"
    ws["B4"] = schema["meta"].get("slice", "")
    ws["A5"] = "Answers source"
    ws["B5"] = str(answers_path)
    ws["A6"] = "Generated at"
    ws["B6"] = datetime.now().isoformat(timespec="seconds")
    ws.sheet_state = "hidden"


# ═══════════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════════

def main():
    if len(sys.argv) not in (3, 4):
        print("Usage: build_model.py <schema.yaml> <answers.yaml> [output.xlsx]")
        sys.exit(2)

    schema_path  = Path(sys.argv[1])
    answers_path = Path(sys.argv[2])
    schema  = yaml.safe_load(schema_path.read_text())
    answers = yaml.safe_load(answers_path.read_text()) or {}

    known = {a["id"] for a in schema_assumptions(schema)}
    unknown = sorted(set(answers) - known)
    if unknown:
        print(f"WARNING: unknown answer keys (ignored): {', '.join(unknown)}",
              file=sys.stderr)

    wb = Workbook()
    wb.remove(wb.active)

    build_assumptions(wb, schema, answers)
    archetypes = schema["meta"].get("archetypes", ["balance"])
    values = {a["id"]: answers.get(a["id"], a.get("default"))
              for a in schema_assumptions(schema)}
    if "subscription" in archetypes:
        import vertical_saas
        vertical_saas.build(wb, schema, values)
        kind = "saas"
    elif "capacity" in archetypes:
        import vertical_power
        vertical_power.build(wb, schema, values)
        kind = "power"
    else:
        warn_milestones(values)
        ri = build_timeline(wb, schema, values)
        build_balance_ramp(wb, ri)
        build_revenue(wb, ri)
        build_opex(wb, ri)
        build_working_capital(wb, ri)
        build_pnl(wb, ri)
        build_forward_cash(wb, ri)
        build_balance_sheet(wb, ri)
        build_cash_flow(wb, ri)
        build_kpi(wb, ri)
        kind = "balance"
    import dashboard
    dashboard.build(wb, kind)
    build_manifest(wb, schema, answers_path)

    if len(sys.argv) == 4:
        out_path = Path(sys.argv[3])
    else:
        out_path = Path(__file__).resolve().parent.parent / "output" / f"{answers_path.stem}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
